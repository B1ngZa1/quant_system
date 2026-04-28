from pptx import Presentation
import sys
import io

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

def extract_ppt_content(pptx_path):
    prs = Presentation(pptx_path)
    content = []
    for i, slide in enumerate(prs.slides, 1):
        slide_text = [f"\n=== 第{i}页 ==="]
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_text.append(shape.text.strip())
        content.append("\n".join(slide_text))
    return "\n".join(content)

if __name__ == "__main__":
    if len(sys.argv) > 1:
        print(extract_ppt_content(sys.argv[1]))
    else:
        print("Usage: python extract_ppt.py <pptx_path>")
