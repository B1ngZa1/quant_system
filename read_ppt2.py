import sys
from pptx import Presentation

prs = Presentation('C:/University/Files/05_学术报告与课程作业/统计方法的应用.pptx')
with open('C:/University/quant_system/ppt2_content.txt', 'w', encoding='utf-8') as f:
    for i, slide in enumerate(prs.slides):
        f.write(f'--- Slide {i+1} ---\n')
        for shape in slide.shapes:
            if hasattr(shape, 'text') and shape.text.strip():
                f.write(shape.text[:400] + '\n')
print("done")
